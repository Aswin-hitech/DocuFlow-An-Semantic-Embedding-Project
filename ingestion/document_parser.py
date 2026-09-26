import os
import io
from typing import Dict, Any, List, Optional, Union

from ingestion.cleaner import clean_text, extract_metadata


class DocumentParser:
    """
    Multi-format document parser supporting PDF, DOCX, TXT, MD, and CSV.
    Extracts structured text and document metadata.
    """

    @staticmethod
    def parse_pdf(file_source: Union[str, bytes, io.BytesIO]) -> Dict[str, Any]:
        """
        Parse text and page metadata from a PDF file path or byte stream.
        Tries PyMuPDF (fitz) first, falls back to pypdf.
        """
        pages_content = []
        metadata = {"file_type": "pdf", "pages": 0}

        # Attempt with PyMuPDF (fitz)
        try:
            import fitz
            if isinstance(file_source, bytes):
                doc = fitz.open(stream=file_source, filetype="pdf")
            elif isinstance(file_source, io.BytesIO):
                doc = fitz.open(stream=file_source.getvalue(), filetype="pdf")
            else:
                doc = fitz.open(file_source)

            metadata["pages"] = len(doc)
            for page_num in range(len(doc)):
                page = doc[page_num]
                text = page.get_text()
                if text.strip():
                    pages_content.append({
                        "page_number": page_num + 1,
                        "text": clean_text(text)
                    })
            doc.close()
        except ImportError:
            # Fallback to pypdf
            import pypdf
            if isinstance(file_source, (bytes, bytearray)):
                stream = io.BytesIO(file_source)
            elif isinstance(file_source, io.BytesIO):
                stream = file_source
            else:
                stream = open(file_source, "rb")

            reader = pypdf.PdfReader(stream)
            metadata["pages"] = len(reader.pages)
            for page_num, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    pages_content.append({
                        "page_number": page_num + 1,
                        "text": clean_text(text)
                    })

            if not isinstance(file_source, (bytes, io.BytesIO)):
                stream.close()

        full_text = "\n\n".join(p["text"] for p in pages_content)
        metadata.update(extract_metadata(full_text))

        return {
            "text": full_text,
            "pages": pages_content,
            "metadata": metadata
        }

    @staticmethod
    def parse_docx(file_source: Union[str, bytes, io.BytesIO]) -> Dict[str, Any]:
        """
        Parse text from a Microsoft Word .docx file.
        """
        import docx

        if isinstance(file_source, bytes):
            stream = io.BytesIO(file_source)
        elif isinstance(file_source, io.BytesIO):
            stream = file_source
        else:
            stream = file_source

        doc = docx.Document(stream)
        paragraphs = []
        for p in doc.paragraphs:
            if p.text.strip():
                paragraphs.append(p.text.strip())

        # Also extract table cells if any
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                row_str = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_str:
                    table_texts.append(row_str)

        all_text = "\n\n".join(paragraphs)
        if table_texts:
            all_text += "\n\n--- Tables ---\n" + "\n".join(table_texts)

        full_text = clean_text(all_text)
        metadata = {"file_type": "docx", "paragraphs_count": len(paragraphs)}
        metadata.update(extract_metadata(full_text))

        return {
            "text": full_text,
            "paragraphs": paragraphs,
            "metadata": metadata
        }

    @staticmethod
    def parse_txt(file_source: Union[str, bytes, io.BytesIO], encoding: str = "utf-8") -> Dict[str, Any]:
        """
        Parse raw plain text, markdown, or code file.
        """
        if isinstance(file_source, bytes):
            text = file_source.decode(encoding, errors="replace")
        elif isinstance(file_source, io.BytesIO):
            text = file_source.getvalue().decode(encoding, errors="replace")
        else:
            with open(file_source, "r", encoding=encoding, errors="replace") as f:
                text = f.read()

        cleaned = clean_text(text)
        metadata = {"file_type": "txt"}
        metadata.update(extract_metadata(cleaned))

        return {
            "text": cleaned,
            "metadata": metadata
        }

    @classmethod
    def parse_file(cls, file_path: str) -> Dict[str, Any]:
        """
        Automatically detect file extension and parse accordingly.
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()
        filename = os.path.basename(file_path)
        file_size = os.path.getsize(file_path)

        if ext == ".pdf":
            result = cls.parse_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            result = cls.parse_docx(file_path)
        elif ext in [".txt", ".md", ".csv", ".json", ".log"]:
            result = cls.parse_txt(file_path)
        else:
            # Fallback to plain text
            result = cls.parse_txt(file_path)

        result["metadata"]["filename"] = filename
        result["metadata"]["file_size_bytes"] = file_size
        result["metadata"]["source_path"] = file_path
        return result

    @classmethod
    def parse_directory(cls, dir_path: str, recursive: bool = True) -> List[Dict[str, Any]]:
        """
        Parse all supported documents in a directory.
        """
        parsed_docs = []
        supported_extensions = {".pdf", ".docx", ".txt", ".md", ".csv", ".json"}

        walker = os.walk(dir_path) if recursive else [(dir_path, [], os.listdir(dir_path))]
        for root, _, files in walker:
            for f in files:
                ext = os.path.splitext(f)[1].lower()
                if ext in supported_extensions:
                    full_path = os.path.join(root, f)
                    try:
                        doc = cls.parse_file(full_path)
                        parsed_docs.append(doc)
                    except Exception as e:
                        print(f"Error parsing file {full_path}: {e}")

        return parsed_docs


if __name__ == "__main__":
    sample_txt_path = "data/raw/documents/sample.txt"
    os.makedirs(os.path.dirname(sample_txt_path), exist_ok=True)
    with open(sample_txt_path, "w", encoding="utf-8") as f:
        f.write("DocuFlow is a cutting-edge semantic embedding framework for enterprise knowledge discovery.")

    parsed = DocumentParser.parse_file(sample_txt_path)
    print("Parsed Result:")
    print("Text:", parsed["text"])
    print("Metadata:", parsed["metadata"])
